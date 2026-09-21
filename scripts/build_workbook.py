#!/usr/bin/env python3
"""Build the four-sheet deliverable with portable openpyxl."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


SHEETS = [
    ("Vehicle Ranking", "vehicle_ranking"),
    ("Listing Matrix", "listing_matrix"),
    ("PLP Matrix", "plp_matrix"),
    ("Negative Keywords", "negative_keywords"),
]


def build_sheet(workbook: Workbook, name: str, rows: list[dict]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    if not rows:
        sheet.append(["No data"])
        return

    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])

    header_fill = PatternFill("solid", fgColor="243447")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="17212B")
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for index, header in enumerate(headers, 1):
        longest = max(len(str(header)), *(len(str(row.get(header, ""))) for row in rows))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 60)

    end = f"{get_column_letter(len(headers))}{len(rows) + 1}"
    table_name = name.replace(" ", "") + "Table"
    table = Table(displayName=table_name, ref=f"A1:{end}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.auto_filter.ref = f"A1:{end}"


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_workbook.py <workbook_data.json> <output.xlsx>")
    input_path, output_path = map(Path, sys.argv[1:])
    payload = json.loads(input_path.read_text(encoding="utf-8"))
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, key in SHEETS:
        build_sheet(workbook, sheet_name, payload.get(key, []))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(json.dumps({"outputPath": str(output_path), "sheets": len(SHEETS)}))


if __name__ == "__main__":
    main()
