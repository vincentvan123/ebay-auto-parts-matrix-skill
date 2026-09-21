import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "build_workbook.py"
SPEC = importlib.util.spec_from_file_location("build_workbook", SCRIPT)
BUILDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILDER)


class WorkbookTests(unittest.TestCase):
    def test_builds_four_expected_sheets(self):
        payload = {
            "vehicle_ranking": [{"SKU": "TEST1", "Rank": 1, "Source": "Hidden", "Source URL": "https://example.test"}],
            "listing_matrix": [{"SKU": "TEST1", "Listing ID": "TEST1-L01"}],
            "plp_matrix": [{"SKU": "TEST1", "Keyword": "sample part"}],
            "negative_keywords": [{"Campaign": "TEST1-L01", "Negative Keyword": "other"}],
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "matrix.xlsx"
            workbook = BUILDER.Workbook()
            workbook.remove(workbook.active)
            for sheet_name, key, table_name in BUILDER.SHEETS:
                BUILDER.build_sheet(workbook, sheet_name, payload[key], table_name)
            workbook.save(output)
            rendered = load_workbook(output, read_only=True)
            self.assertEqual(rendered.sheetnames, [name for name, _, _ in BUILDER.SHEETS])
            self.assertEqual(rendered["Listing 矩阵"]["B2"].value, "TEST1-L01")
            ranking_headers = [cell.value for cell in rendered["车型市场排名"][1]]
            self.assertIn("排名", ranking_headers)
            self.assertNotIn("Source", ranking_headers)
            self.assertNotIn("Source URL", ranking_headers)


if __name__ == "__main__":
    unittest.main()
